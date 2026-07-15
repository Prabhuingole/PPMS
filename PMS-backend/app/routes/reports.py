import io
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.orm import Session
from sqlalchemy import func, and_, or_, case, literal_column
from typing import List, Optional
from datetime import date, datetime, timedelta, timezone
import calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.database.session import get_db
from app.core.deps import get_current_user
from app.models.user import User
from app.models.attendance import Attendance
from app.models.timesheet import Timesheet
from app.models.task import Task
from app.models.project import Project
from app.models.company import Company
from app.schemas.reports import (
    AttendanceSummaryRecord,
    DailyAttendanceRecord,
    TimesheetComplianceRecord,
    ProductivityRecord,
    MissingAttendanceRecord
)

router = APIRouter(prefix="/reports", tags=["Reports"])

def apply_tenant_filter(query, model, current_user: User):
    if current_user.role in ["superadmin"]:
        return query
    elif current_user.role in ["admin"]:
        # Assuming model has company_id or belongs to user with company_id
        if hasattr(model, "company_id"):
            return query.filter(model.company_id == current_user.company_id)
        elif hasattr(model, "user_id"):
            return query.join(User, model.user_id == User.id).filter(User.company_id == current_user.company_id)
        else:
            return query
    else:
        raise HTTPException(status_code=403, detail="Not authorized to access reports")

def apply_common_filters(query, model, company_id, project_id, role, user_id, start_date, end_date):
    if hasattr(model, "date"):
        if start_date: query = query.filter(model.date >= start_date)
        if end_date: query = query.filter(model.date <= end_date)
    elif hasattr(model, "created_at"):
        if start_date: query = query.filter(model.created_at >= start_date)
        if end_date: query = query.filter(model.created_at <= end_date)
        
    if user_id and hasattr(model, "user_id"):
        query = query.filter(model.user_id == user_id)
        
    return query

@router.get("/attendance", response_model=List[AttendanceSummaryRecord])
def get_attendance_summary(
    company_id: Optional[int] = None,
    project_id: Optional[int] = None,

    role: Optional[str] = None,
    user_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role not in ["superadmin", "admin"]:
        raise HTTPException(status_code=403, detail="Not authorized")
        
    # Baseline users query
    users_q = db.query(User).filter(User.role.in_(["employee", "intern"]))
    users_q = apply_tenant_filter(users_q, User, current_user)
    
    if company_id: users_q = users_q.filter(User.company_id == company_id)

    if role: users_q = users_q.filter(User.role == role)
    if user_id: users_q = users_q.filter(User.id == user_id)
    
    users = users_q.all()
    results = []
    
    if not start_date:
        start_date = date.today() - timedelta(days=30)
    if not end_date:
        end_date = date.today()
        
    # Total days calculation (excluding Sundays)
    total_days = 0
    current = start_date
    while current <= end_date:
        if current.weekday() != 6: # 6 is Sunday
            total_days += 1
        current += timedelta(days=1)
        
    for u in users:
        # Fetch attendance records for this user in date range
        att_records = db.query(Attendance).filter(
            Attendance.user_id == u.id,
            Attendance.date >= start_date,
            Attendance.date <= end_date
        ).all()
        
        present_days = 0
        missing_in = 0
        missing_out = 0
        total_hours = 0.0
        last_clock_in = None
        last_clock_out = None
        
        for att in att_records:
            if att.date.weekday() == 6:
                continue
            if att.clock_in_time: last_clock_in = att.clock_in_time
            if att.clock_out_time: last_clock_out = att.clock_out_time
            if att.clock_in_time and att.clock_out_time:
                present_days += 1
            elif att.clock_in_time and not att.clock_out_time:
                missing_out += 1
            elif att.clock_out_time and not att.clock_in_time:
                missing_in += 1
                
            if att.working_hours:
                total_hours += float(att.working_hours)
                
        absent_days = max(0, total_days - present_days - missing_in - missing_out)
        avg_hours = total_hours / present_days if present_days > 0 else 0
        
        company_name = db.query(Company.name).filter(Company.id == u.company_id).scalar() or "Unknown"

        
        results.append(AttendanceSummaryRecord(
            user_id=u.id,
            name=u.full_name or "Unknown",
            employee_id=u.employee_id or str(u.id),
            company=company_name,
            role=u.role,

            total_working_days=total_days,
            present_days=present_days,
            absent_days=absent_days,
            half_days=0,
            leave_days=0,
            missing_clock_in=missing_in,
            missing_clock_out=missing_out,
            late_arrivals=0,
            early_leaves=0,
            average_working_hours=round(avg_hours, 2),
            total_working_hours=round(total_hours, 2),
            overtime_hours=0.0,
            clock_in=last_clock_in,
            clock_out=last_clock_out
        ))
        
    return results

@router.get("/attendance/daily", response_model=List[DailyAttendanceRecord])
def get_daily_attendance(
    company_id: Optional[int] = None,
    user_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role not in ["superadmin", "admin"]:
        raise HTTPException(status_code=403, detail="Not authorized")
        
    if not start_date:
        start_date = date.today() - timedelta(days=7)
    if not end_date:
        end_date = date.today()
        
    # Baseline users query
    users_q = db.query(User).filter(User.role.in_(["employee", "intern"]))
    users_q = apply_tenant_filter(users_q, User, current_user)
    
    if company_id: users_q = users_q.filter(User.company_id == company_id)
    if user_id: users_q = users_q.filter(User.id == user_id)
    
    users = users_q.all()
    user_ids = [u.id for u in users]
    
    # Fetch attendances
    attendances = db.query(Attendance).filter(
        Attendance.user_id.in_(user_ids),
        Attendance.date >= start_date,
        Attendance.date <= end_date
    ).all()
    
    # Fetch timesheets
    timesheets = db.query(Timesheet.user_id, Timesheet.date).filter(
        Timesheet.user_id.in_(user_ids),
        Timesheet.date >= start_date,
        Timesheet.date <= end_date
    ).distinct().all()
    
    # Create lookups
    att_map = {(a.user_id, a.date): a for a in attendances}
    ts_map = set((t.user_id, t.date.date() if hasattr(t.date, 'date') else t.date) for t in timesheets)
    
    results = []
    
    # Generate dates excluding Sundays
    current = start_date
    dates_to_check = []
    while current <= end_date:
        if current.weekday() != 6:
            dates_to_check.append(current)
        current += timedelta(days=1)
        
    for u in users:
        company_name = u.company.name if u.company else "Unknown"
        
        for d in dates_to_check:
            att = att_map.get((u.id, d))
            ts_filled = (u.id, d) in ts_map
            
            status = "Absent"
            clock_in = None
            clock_out = None
            total_hours = 0.0
            
            if att:
                clock_in = att.clock_in_time
                clock_out = att.clock_out_time
                total_hours = att.working_hours or 0.0
                
                if clock_in and clock_out:
                    status = "Present"
                elif clock_in and not clock_out:
                    now_ist = datetime.now(timezone.utc) + timedelta(hours=5, minutes=30)
                    # If it's today and time is before 7 PM (19:00), status is "Yet to Clock Out"
                    if d == now_ist.date() and now_ist.hour < 19:
                        status = "Yet to Clock Out"
                    else:
                        status = "Missing Clock Out"
                elif clock_out and not clock_in:
                    status = "Missing Clock In"
                    
            results.append(DailyAttendanceRecord(
                user_id=u.id,
                name=u.full_name or "Unknown",
                employee_id=u.employee_id or str(u.id),
                company=company_name,
                role=u.role,
                date=d,
                clock_in=clock_in,
                clock_out=clock_out,
                total_hours=round(total_hours, 2),
                status=status,
                timesheet_filled=ts_filled
            ))
            
    # Sort results by date desc, then name
    results.sort(key=lambda x: (x.date, x.name), reverse=True)
    
    return results

@router.get("/timesheet-compliance", response_model=List[TimesheetComplianceRecord])
def get_timesheet_compliance(
    company_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # Simplified placeholder for complex aggregations
    return []

@router.get("/productivity", response_model=List[ProductivityRecord])
def get_productivity(
    company_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # Simplified placeholder for complex aggregations
    return []

@router.get("/missing-attendance", response_model=List[MissingAttendanceRecord])
def get_missing_attendance(
    company_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # Simplified placeholder for complex aggregations
    return []

@router.post("/export")
def export_excel(
    report_type: str = Query(..., description="attendance | daily | timesheet | productivity | missing"),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role not in ["superadmin", "admin"]:
        raise HTTPException(status_code=403, detail="Not authorized")
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    
    # Title Row
    ws.append([f"SPMS {report_type.capitalize()} Report"])
    ws.append([f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws.append([])
    
    if report_type == "attendance":
        headers = ["Employee ID", "Name", "Company", "Role", "Present Days", "Absent Days", "Missing Clock Out", "Total Hours", "Avg Hours", "Last Clock In", "Last Clock Out"]
        ws.append(headers)
        
        # Style headers
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
            
        data = get_attendance_summary(company_id=company_id, start_date=start_date, end_date=end_date, db=db, current_user=current_user)
        for r in data:
            ws.append([
                r.employee_id, r.name, r.company, r.role,
                r.present_days, r.absent_days, r.missing_clock_out,
                r.total_working_hours, r.average_working_hours,
                (r.clock_in + timedelta(hours=5, minutes=30)).strftime('%Y-%m-%d %I:%M %p') if r.clock_in else "N/A",
                (r.clock_out + timedelta(hours=5, minutes=30)).strftime('%Y-%m-%d %I:%M %p') if r.clock_out else "N/A"
            ])
            
    elif report_type == "daily":
        headers = ["Date", "Employee ID", "Name", "Company", "Role", "Clock In", "Clock Out", "Total Hours", "Status", "Timesheet Filled"]
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
            
        data = get_daily_attendance(company_id=company_id, start_date=start_date, end_date=end_date, db=db, current_user=current_user)
        for r in data:
            ws.append([
                r.date.strftime('%Y-%m-%d'),
                r.employee_id, r.name, r.company, r.role,
                (r.clock_in + timedelta(hours=5, minutes=30)).strftime('%I:%M %p') if r.clock_in else "—",
                (r.clock_out + timedelta(hours=5, minutes=30)).strftime('%I:%M %p') if r.clock_out else "—",
                r.total_hours, r.status,
                "Yes" if r.timesheet_filled else "No"
            ])
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    headers = {
        'Content-Disposition': f'attachment; filename="Report_{report_type}_{date.today()}.xlsx"'
    }
    
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )
