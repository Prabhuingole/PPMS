import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def create_spms_presentation():
    # Create presentation
    prs = Presentation()
    
    # Define slide layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    
    # -----------------------------------------------------------------
    # Slide 1: Title Slide
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Sambhav Project Management System (SPMS)"
    subtitle.text = "Enterprise-Grade Project & Resource Management\n\nConfidential & Proprietary"

    # -----------------------------------------------------------------
    # Slide 2: Motive & Vision
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Our Motive & Vision"
    tf = body_shape.text_frame
    
    tf.text = "Why SPMS was created:"
    
    p = tf.add_paragraph()
    p.text = "Unified Hub: Eliminate fragmented tools by centralizing project tracking, timesheets, and documentation."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Security First: Implement robust enterprise-grade security including Multi-Factor Authentication (MFA) and immutable audit logs."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Scalable Architecture: Built to handle organizational growth with flexible Role-Based Access Control (RBAC)."
    p.level = 1

    # -----------------------------------------------------------------
    # Slide 3: Core Features
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Core Features"
    tf = body_shape.text_frame
    
    tf.text = "Comprehensive feature suite:"
    
    features = [
        "Multi-Superadmin Management: Provision distinct, fully trackable superadmin/admin profiles securely.",
        "Automated Timesheets: Frictionless time logging tied to specific project modules.",
        "Document Templates & Storage: Centralized generation and storage of critical project documents.",
        "Real-Time Notifications: Powered by WebSockets to keep distributed teams instantly synced.",
        "Rich Analytics & Logs: Monitor granular user activity through a centralized Activity Logs dashboard."
    ]
    
    for feature in features:
        p = tf.add_paragraph()
        p.text = feature
        p.level = 1

    # -----------------------------------------------------------------
    # Slide 4: Usage & Role Perspectives
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Usage & Role Perspectives"
    tf = body_shape.text_frame
    
    tf.text = "Tailored experiences for every level of the organization:"
    
    roles = [
        ("Employees / Interns", "Focused dashboard for submitting timesheets, viewing assigned projects, and accessing shared templates."),
        ("Reporting Managers", "Dedicated interface to review team performance, approve timesheets, and manage project milestones."),
        ("Superadmins / Admins", "Global overview, system configuration, deep audit access, and provisioning of new root-level accounts.")
    ]
    
    for role, desc in roles:
        p = tf.add_paragraph()
        p.text = role
        p.level = 1
        p.font.bold = True
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.level = 2

    # -----------------------------------------------------------------
    # Slide 5: Project Details & Tech Stack
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Tech Stack & Architecture"
    tf = body_shape.text_frame
    
    tf.text = "Built on modern, high-performance technologies:"
    
    stack = [
        "Frontend: Next.js 15, React 19, TailwindCSS, and Framer Motion for beautiful, fluid interfaces.",
        "Backend: FastAPI (Python) for blazing-fast asynchronous processing and API delivery.",
        "Database: SQLAlchemy ORM handling relational integrity with robust migration support.",
        "Security: JWT-based stateless authentication, bcrypt password hashing, and role-based middleware."
    ]
    
    for tech in stack:
        p = tf.add_paragraph()
        p.text = tech
        p.level = 1

    # -----------------------------------------------------------------
    # Slide 6: Summary & Impact
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Impact & Next Steps"
    tf = body_shape.text_frame
    
    tf.text = "Sambhav PMS is positioned to:"
    
    impacts = [
        "Drastically reduce administrative overhead.",
        "Provide 100% transparency into project health and employee performance.",
        "Secure sensitive company data under a strict zero-trust philosophy."
    ]
    
    for impact in impacts:
        p = tf.add_paragraph()
        p.text = impact
        p.level = 1
        
    p_final = tf.add_paragraph()
    p_final.text = "\nThank You."
    p_final.font.bold = True

    # Save presentation
    output_path = os.path.join(os.path.dirname(__file__), "SPMS_Overview_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == '__main__':
    create_spms_presentation()
